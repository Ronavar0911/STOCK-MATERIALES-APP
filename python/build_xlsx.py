import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

F = 'Arial'
H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F4E5A')
N = Font(name=F, size=10); B = Font(name=F, size=10, bold=True); IN = Font(name=F, size=10, color='0000FF')
Y = PatternFill('solid', fgColor='FFF2CC')
base = pd.read_pickle('base.pkl')
wb = Workbook()

def table(ws, df, ref_row, name, style='TableStyleMedium2', widths=None):
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(ref_row, j, c); cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='center')
    for i, row in enumerate(df.itertuples(index=False), ref_row + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and np.isnan(v): v = None
            if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
            ws.cell(i, j, v).font = N
    last = ref_row + max(len(df), 1)
    t = Table(displayName=name, ref=f'A{ref_row}:{get_column_letter(len(df.columns))}{last}')
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(12, min(40, len(c) + 4)))

# ---------- LEEME ----------
ws = wb.active; ws.title = 'LEEME'
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 120
lines = [
 ('t', 'APP STOCK DE MATERIALES – Motor Power Query'),
 ('', 'Lee el archivo .xlsx más reciente de cada carpeta SAP y calcula stock vs mínimo, seguimiento de SOLPED y consumo de la temporada.'),
 ('h', 'Cómo se actualiza'),
 ('', '1. Guarda el nuevo export de SAP en su carpeta (no hace falta borrar el anterior: se toma el más reciente por fecha de modificación).'),
 ('', '2. En este libro: Datos > Actualizar todo. Se recalculan TABLERO, SOLPED_SEGUIMIENTO, MOV_TEMPORADA, RESERVAS_OT y OTS.'),
 ('', '3. Sube este libro a la app web (botón "Cargar Excel actualizado") para que el equipo vea los datos nuevos.'),
 ('h', 'Configuración inicial (una sola vez, ~10 min)'),
 ('', '1. Hoja PARAMETROS: revisa las rutas de tblCarpetas (MB52 en …\\01_SAP_EXPORTS\\MATERIALES_STOCK).'),
 ('', '2. Datos > Obtener datos > Desde otras fuentes > Consulta en blanco. Abre "Editor avanzado", pega el código y renombra la consulta EXACTAMENTE:'),
 ('', '     SAP  →  Cerrar y cargar en… > Solo crear conexión'),
 ('', '     SOLPED_SEGUIMIENTO, MOV_TEMPORADA, RESERVAS_OT, TABLERO, OTS  →  Cerrar y cargar en… > Tabla en hoja nueva'),
 ('', '   El código de cada consulta está en el archivo CONSULTAS_POWER_QUERY.txt (en ese orden).'),
 ('', '3. Si Excel pregunta por niveles de privacidad: Archivo > Opciones > Privacidad > "Ignorar niveles de privacidad" (todo es local).'),
 ('h', 'Qué exportar de SAP'),
 ('', 'MB51: movimientos de los ÚLTIMOS 12 MESES (todas las clases de movimiento). Así se ven las recepciones de OC antiguas y el consumo de la temporada.'),
 ('', '      Recomendado: añadir la columna "Orden" al layout de MB51 → el consumo se clasifica solo en OM01 / OM03 / OM04.'),
 ('', 'MB52: stock actual, todos los almacenes del centro.   ME5A: SOLPED (con pedido y cantidad pedida).'),
 ('', 'IW13: utilización de material en OTs.   IW39: OTs.   IH08: equipos de riego.'),
 ('h', 'Reglas de cálculo'),
 ('', 'Temporada: de la semana ISO 27 a la semana 26 del año siguiente (temporada actual 2026/2027 inicia el 29/06/2026). Se calcula sola.'),
 ('', 'Consumo = mov. 201 + 261 + 221 menos sus anulaciones 202 + 262 + 222. Entradas = 101 − 102. Traslados (311, 309…) solo se muestran.'),
 ('', 'Almacenes: todos excepto los de AlmacenesExcluidos (1030).'),
 ('', 'BASE_HISTORICA (estática): temporadas completas 2023/24, 2024/25 y 2025/26 = 36 meses. Se recalcula una vez por temporada.'),
 ('', '  · Promedio mensual ajustado: los meses con consumo mayor a 3 × la mediana de los meses con consumo se recortan a ese tope (evita que un proyecto infle el promedio).'),
 ('', '  · Tipo de demanda: Regular (≥18 meses con consumo), Intermitente (6–17), Esporádico (<6), Puntual (consumo en una sola temporada), Inactivo (sin consumo en 12 meses).'),
 ('', 'Mínimo = MAX( Prom. ajustado × MesesMin  [solo Regular e Intermitente] ;  Mínimo manual de CRITICOS ).   Máximo = Prom. ajustado × MesesMax.'),
 ('', 'Estado: Quiebre (stock 0 y nada pedido) · Pedir (stock + SOLPED + OC < mínimo) · En camino (stock < mínimo pero lo pedido lo cubre) · OK · Sobrestock (> Máximo × FactorSobrestock) · Sin mínimo.'),
 ('', 'Cantidad sugerida a pedir = Máximo − (stock + OC en tránsito + SOLPED sin OC), solo cuando está por debajo del mínimo.'),
 ('', 'SOLPED: 1 Solicitado sin OC · 2 OC en tránsito · 3 Recibido parcial · 4 Recibido (MB51 101 del mismo pedido y material) · 5 Servicio con OC.'),
 ('h', 'Celdas que se editan'),
 ('', 'Texto azul sobre fondo amarillo = dato editable (PARAMETROS y CRITICOS). Todo lo demás lo calcula Power Query.'),
]
r = 1
for k, txt in lines:
    c = ws.cell(r, 2, txt)
    c.font = Font(name=F, size=16, bold=True, color='1F4E5A') if k == 't' else Font(name=F, size=11, bold=True, color='1F4E5A') if k == 'h' else N
    if k == 'h': r += 0
    c.alignment = Alignment(wrap_text=True, vertical='top')
    r += 2 if k == 'h' and False else 1
    if k in ('t',): r += 1

# ---------- PARAMETROS ----------
ws = wb.create_sheet('PARAMETROS')
par = pd.DataFrame([
    ('MesesMin', 2, 'Meses de consumo ajustado que cubre el stock mínimo (tiempo de reposición + seguridad).'),
    ('MesesMax', 4, 'Meses de consumo ajustado para el stock máximo / cantidad a pedir.'),
    ('FactorSobrestock', 1.5, 'Stock > Máximo × este factor se marca como Sobrestock.'),
    ('AlmacenesExcluidos', '1030', 'Almacenes separados por coma que no se consideran.'),
], columns=['Parametro', 'Valor', 'Descripcion'])
ws['A1'] = 'Parámetros de cálculo'; ws['A1'].font = Font(name=F, size=13, bold=True, color='1F4E5A')
table(ws, par, 3, 'tblParametros', widths={'Parametro': 22, 'Valor': 14, 'Descripcion': 80})
for i in range(4, 8):
    ws.cell(i, 2).font = IN; ws.cell(i, 2).fill = Y
root = r'C:\Users\srosales\Hortifrut Chile SA\HFBP-AGR-PMA - Documentos\PM - C&G T2627\PM - RIEGO\MAESTRO_MANTENIMIENTO_RIEGO\01_SAP_EXPORTS'
carp = pd.DataFrame([('MB51', root + r'\CONSUMOS'), ('MB52', root + r'\MATERIALES_STOCK'), ('ME5A', root + r'\SOLPED'),
                     ('IW13', root + r'\UTILIZACION_MATERIAL'), ('IW39', root + r'\OT'), ('IH08', root + r'\EQUIPOS')],
                    columns=['Fuente', 'Carpeta'])
ws['A10'] = 'Carpetas de exports SAP (se toma el .xlsx más reciente de cada una)'; ws['A10'].font = Font(name=F, size=13, bold=True, color='1F4E5A')
for j, c in enumerate(carp.columns, 1):
    x = ws.cell(12, j, c); x.font = H; x.fill = HF
for i, row in enumerate(carp.itertuples(index=False), 13):
    ws.cell(i, 1, row[0]).font = N
    c = ws.cell(i, 2, row[1]); c.font = IN; c.fill = Y
t = Table(displayName='tblCarpetas', ref='A12:B18'); t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True); ws.add_table(t)

ws.column_dimensions['B'].width = 120

# ---------- CRITICOS ----------
ws = wb.create_sheet('CRITICOS')
ws['A1'] = 'Materiales críticos – mínimo manual'; ws['A1'].font = Font(name=F, size=13, bold=True, color='1F4E5A')
ws['A2'] = 'Completa Material y Minimo_manual (texto azul). El mínimo final será el MAYOR entre este valor y el sugerido por consumo. La fila EJEMPLO no afecta los cálculos; puedes borrarla.'
ws['A2'].font = Font(name=F, size=10, italic=True)
cr = pd.DataFrame([('EJEMPLO-001', None, 10, 'Repuesto de bomba sin reemplazo inmediato', 'S. Rosales', pd.Timestamp('2026-10-04'))],
                  columns=['Material', 'Descripcion', 'Minimo_manual', 'Motivo', 'Responsable', 'Fecha'])
table(ws, cr, 4, 'tblCriticos', widths={'Material': 16, 'Descripcion': 45, 'Minimo_manual': 16, 'Motivo': 50, 'Responsable': 18, 'Fecha': 14})
for rr in range(5, 6):
    for cc in (1, 3, 4, 5, 6):
        ws.cell(rr, cc).font = IN; ws.cell(rr, cc).fill = Y
    ws.cell(rr, 2, f'=IFERROR(INDEX(BASE_HISTORICA!$B:$B,MATCH(A{rr},BASE_HISTORICA!$A:$A,0)),"")').font = N
    ws.cell(rr, 6).number_format = 'dd/mm/yyyy'

# ---------- BASE_HISTORICA ----------
ws = wb.create_sheet('BASE_HISTORICA')
b = base.copy()
table(ws, b, 1, 'tblBaseHistorica', widths={'Descripcion': 42, 'Material': 14})
fmt = {'Prom_mensual_simple': '#,##0.00', 'Prom_mensual_ajustado': '#,##0.00', 'Pico_pct': '0%', 'Pct_proyecto': '0%',
       'Ultimo_consumo': 'dd/mm/yyyy', 'Cons_2023_2024': '#,##0.##', 'Cons_2024_2025': '#,##0.##', 'Cons_2025_2026': '#,##0.##',
       'Cons_total': '#,##0.##', 'Mes_pico': '#,##0.##'}
for j, c in enumerate(b.columns, 1):
    if c in fmt:
        for i in range(2, len(b) + 2):
            ws.cell(i, j).number_format = fmt[c]
ws.freeze_panes = 'C2'
notes = {'Prom_mensual_ajustado': 'Promedio de 36 meses con los picos (> 3× mediana de meses con consumo) recortados.',
         'Pico_pct': '% del consumo total concentrado en el mes de mayor consumo.',
         'Pct_proyecto': '% del consumo que fue salida a proyecto (mov. 221).',
         'Temporadas_con_consumo': 'De 3 temporadas completas (2023/24 a 2025/26).'}
for j, c in enumerate(b.columns, 1):
    if c in notes: ws.cell(1, j).comment = Comment(notes[c], 'Claude')

wb.save('/mnt/user-data/outputs/APP_STOCK_MATERIALES.xlsx')
print('ok', len(b))
